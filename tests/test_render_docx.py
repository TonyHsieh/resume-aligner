"""Golden tests: the renderer must reproduce SAMPLE_FORMAT.docx's formatting."""

from pathlib import Path

import docx
import pytest
from docx.oxml.ns import qn

from resume_factory.model import Labeled, Resume, Role, normalize
from resume_factory.parse_docx import _runs, parse_docx
from resume_factory.render_docx import DEFAULT_TEMPLATE, render_resume

SAMPLE = Path(DEFAULT_TEMPLATE)


def _mark(p):
    """(bold, size_halfpts) of the paragraph-mark run properties.

    The mark's italic flag is deliberately ignored: it is invisible (it only
    styles the pilcrow) and the sample itself is inconsistent about it.
    """
    ppr = p._p.pPr
    rpr = ppr.find(qn("w:rPr")) if ppr is not None else None
    if rpr is None:
        return None
    sz = rpr.find(qn("w:sz"))
    return (
        rpr.find(qn("w:b")) is not None,
        sz.get(qn("w:val")) if sz is not None else None,
    )


def signature(path):
    """Formatting fingerprint per paragraph, insensitive to run fragmentation."""
    sigs = []
    for p in docx.Document(str(path)).paragraphs:
        ppr = p._p.pPr
        num = ppr.find(qn("w:numPr")) if ppr is not None else None
        # The sample gives every role its own numId, but all 42 bullet
        # definitions are identical (same glyph and indent), so only compare
        # whether the paragraph is a bullet, not which list instance it uses.
        is_bullet = num is not None
        sigs.append(
            (
                p.style.name,
                str(p.alignment),
                ppr is not None and ppr.find(qn("w:pBdr")) is not None,
                is_bullet,
                _mark(p) if not p.text.strip() else None,
                tuple((normalize(t), b, i, s) for t, b, i, s in _runs(p)),
            )
        )
    return sigs


@pytest.fixture(scope="module")
def sample_resume():
    return parse_docx(SAMPLE)


def test_parser_finds_expected_structure(sample_resume):
    r = sample_resume
    assert r.name == "Tony Hsieh"
    assert len(r.roles) == 11
    assert r.roles[0].company == "Stellarus"
    assert all(role.bullets for role in r.roles)
    assert [e.label for e in r.education] == ["University of California, Santa Barbara:"]
    assert len(r.skills) == 3


def test_roundtrip_matches_sample_formatting(sample_resume, tmp_path):
    out = render_resume(sample_resume, tmp_path / "roundtrip.docx")
    original, rendered = signature(SAMPLE), signature(out)
    assert len(rendered) == len(original)
    for i, (a, b) in enumerate(zip(original, rendered)):
        assert a == b, f"paragraph {i} differs:\n sample  : {a}\n rendered: {b}"


def test_page_setup_preserved(sample_resume, tmp_path):
    out = render_resume(sample_resume, tmp_path / "r.docx")
    s = docx.Document(str(out)).sections[0]
    ref = docx.Document(str(SAMPLE)).sections[0]
    assert (s.left_margin, s.right_margin, s.top_margin, s.bottom_margin) == (
        ref.left_margin,
        ref.right_margin,
        ref.top_margin,
        ref.bottom_margin,
    )
    assert (s.page_width, s.page_height) == (ref.page_width, ref.page_height)


def _tiny():
    return Resume(
        name="Test Person",
        headline="Headline with **new keywords**",
        contact="(555) 555-5555 • a@b.com",
        summary="Summary with **changed words** inside.",
        expertise=["A", "B **C**"],
        roles=[
            Role(
                "Acme",
                "2020 – Present",
                "Senior **Platform** Manager",
                [Labeled("Lead:", "Ran **Kubernetes** platform.")],
            )
        ],
        education=[Labeled("School:", "Degree")],
        skills=[Labeled("Tools:", "Python")],
    )


def test_markers_stripped_in_clean_and_bold_in_review(tmp_path):
    clean = render_resume(_tiny(), tmp_path / "clean.docx")
    review = render_resume(_tiny(), tmp_path / "review.docx", review=True)

    clean_text = "\n".join(p.text for p in docx.Document(str(clean)).paragraphs)
    assert "**" not in clean_text
    assert "Ran Kubernetes platform." in clean_text

    def bold_text(path):
        return {
            r.text
            for p in docx.Document(str(path)).paragraphs
            for r in p.runs
            if r.bold and r.font.size and r.font.size.pt == 9
        }

    assert "Kubernetes" not in bold_text(clean)
    assert {"Kubernetes", "changed words", "Platform", "C"} <= bold_text(review)
